# -*- coding: utf-8 -*-
"""项目总结 PPT 生成脚本：运行后生成 项目总结.pptx（16:9）"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

FONT = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)   # 主题蓝
DARK = RGBColor(0x24, 0x29, 0x2F)
GRAY = RGBColor(0x5F, 0x6B, 0x7A)


def set_font(run, size=18, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.makeelement(qn("a:ea"), {"typeface": FONT})
    rpr.append(ea)


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_font(run, size=30, bold=True, color=DARK)
    # 标题下划线装饰条
    bar = slide.shapes.add_shape(1, Inches(0.65), Inches(1.15), Inches(1.6), Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_bullets(slide, items, top=1.55, left=0.7, width=12.0, height=5.4):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for text, level in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = text
        if level == 0:
            set_font(run, size=19, bold=True, color=DARK)
        else:
            set_font(run, size=16, color=GRAY)
    return box


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# ---------- S1 封面 ----------
s = prs.slides.add_slide(blank)
tb = s.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(11), Inches(2.6))
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = "基于 Hadoop + Hive 的电商数仓与可视化平台"; set_font(r, 38, True, DARK)
p2 = tb.text_frame.add_paragraph()
r = p2.add_run(); r.text = "从搭建 WSL2 集群到数仓落地的完整实战之旅"; set_font(r, 22, False, ACCENT)
p3 = tb.text_frame.add_paragraph()
r = p3.add_run(); r.text = "技术栈：Hadoop 3.3.6 / Hive 3.1.3 / Spark 3.5.1 / Superset / PyHive"; set_font(r, 15, False, GRAY)

slides = [
    ("全程路线图", [
        ("① 集群搭建：WSL2 + Hadoop + Hive + Spark 单节点", 0),
        ("② 数据入仓：预处理 → HDFS → ODS → DWD → DWM → DWS → DM", 0),
        ("③ 可视化双路线：Superset 看板 + Python 报表", 0),
        ("④ 进阶三站：Spark 提速 / 定时自动化 / 维度建模", 0),
        ("⑤ 收官：项目文档化、简历沉淀", 0),
    ]),
    ("环境与组件", [
        ("硬件与系统", 0),
        ("Windows 11 + WSL2 Ubuntu 22.04，单节点 8GB 内存（镜像网络模式）", 1),
        ("大数据组件", 0),
        ("Hadoop 3.3.6：HDFS 存储 + YARN 调度 + MapReduce 计算", 1),
        ("Hive 3.1.3：数仓引擎，MySQL 8.0 存元数据，HiveServer2 对外服务", 1),
        ("Spark 3.5.1：Standalone 模式，进阶期主力计算引擎", 1),
        ("应用层", 0),
        ("Superset 6.1.0（BI 看板） / PyHive + pyecharts（Python 报表）", 1),
        ("三种客户端同一入口：localhost:10000（Thrift/SASL 协议）", 1),
    ]),
    ("阶段一：集群搭建 —— 做了什么", [
        ("搭建动作", 0),
        ("SSH 免密、Java 环境、Hadoop 格式化 NameNode 并启动 HDFS/YARN", 1),
        ("Hive 对接 MySQL Metastore，初始化 Schema，启动 Metastore + HiveServer2", 1),
        ("Spark Master/Worker 启动，打通 WSL 与 Windows 双侧访问", 1),
        ("遇到的问题与解法", 0),
        ("Hive 3.1.3 与高版本 Java 不兼容 → 锁定 Java 8 运行 Hive", 1),
        ("Hive 与 Hadoop 的 Guava 版本冲突 → 替换 Hive 内 jar", 1),
        ("MySQL 8.0 认证插件不被识别 → 改用兼容认证方式", 1),
        ("bashrc 重复追加环境变量导致 PATH 过长故障 → 去重清理", 1),
    ]),
    ("阶段二：数据入仓与五层数仓 —— 做了什么", [
        ("数据链路", 0),
        ("淘宝订单 99,457 行：sed 去尾逗号 → cut 裁异常列 → hdfs dfs -put 入仓", 1),
        ("ODS 原样贴源 → DWD 八步清洗（ORC）→ DWM 用户日汇总 → DWS 品类月统计 → DM 报表直供", 1),
        ("遇到的问题与解法", 0),
        ("1 行 14 列脏数据 → awk/uniq 侦查定位，决策裁剪而非删行（保订单）", 1),
        ("日期 2021/4/17 与浮点噪声 60.59999999 → 清洗前先做 6 项指标数据体检", 1),
        ("DM 层遵循华为规范命名（对标阿里 ADS）", 1),
        ("掌握的知识", 0),
        ("数仓分层职责、外部表与跳表头、窗口函数、聚合不可逆（COUNT DISTINCT 不能二次汇总）", 1),
    ]),
    ("阶段三：可视化双路线 —— 做了什么", [
        ("B 路线：Superset 看板", 0),
        ("虚拟环境安装 → 初始化元数据 → 接通 Hive（SQLAlchemy URI）→ 建数据集与双图看板", 1),
        ("遇到的问题：6.1.0 漏依赖（rich / cachetools / thrift_sasl）看 Traceback 逐个补", 1),
        ("图表 60 秒超时被掐 → 外部配置文件双参数放宽至 300 秒", 1),
        ("A 路线：PyHive + pyecharts", 0),
        ("Python 直连数仓取数 → 折线+柱状双图交互报告 → 文件名带日期戳存档", 1),
        ("遇到的问题：库名假设错误（ads ≠ 实际 dm）→ 写 SQL 前先 SHOW DATABASES 验证", 1),
        ("两条路线的取舍：BI 实时但慢，代码报表秒开但无交互——架构权衡题答案", 1),
    ]),
    ("阶段四：进阶三站 —— 做了什么", [
        ("站一：Spark 提速", 0),
        ("PySpark DataFrame API 复刻 DWD 清洗：1.5 秒完成，对比 Hive MR 的 7 分钟（快约 280 倍）", 1),
        ("掌握惰性执行、WholeStageCodegen、谓词下推、列裁剪；Parquet 落盘验证", 1),
        ("站二：定时自动化", 0),
        ("Windows 任务计划程序挂每日报表任务；要点：必须填“起始于”工作目录", 1),
        ("建立意识：先手动验证后挂定时；定时任务要检查外部服务依赖", 1),
        ("站三：维度建模", 0),
        ("星型模型：品类维表（ROW_NUMBER 代理键）+ 客户维表（CASE WHEN 派生年龄段）", 1),
        ("三表 JOIN 查询发现：50+ 人群才是消费榜首，推翻“年轻人主力”的直觉", 1),
    ]),
    ("踩坑记录：环境与连接", [
        ("百度网盘抢占 10000 端口", 0),
        ("镜像模式共享端口空间，HiveServer2 起不来 → 结束占用进程后重启", 1),
        ("JDBC SASL 握手失败", 0),
        ("WSL2 localhost 端口转发代理破坏 SASL 流 → 切镜像模式直连", 1),
        ("Python 环境缺件", 0),
        ("Ubuntu 拆包：缺 python3.10-venv 建不了虚拟环境；缺 python3.10-dev 编译 sasl 失败", 1),
        ("共同方法论：报错先分清客户端还是服务端，再找日志现场", 0),
    ]),
    ("踩坑记录：计算引擎", [
        ("return code 2（第一次）：本地模式 OOM", 0),
        ("local.auto=true 把任务塞进 HiveServer2 的 256MB 小堆 → 删配置 + 堆提到 1GB", 1),
        ("return code 2（第二次）：MRAppMaster 找不到类", 0),
        ("tarball 安装缺 mapreduce.application.classpath → 写死绝对路径类路径", 1),
        ("INSERT 中途 Connection reset", 0),
        ("客户端断连导致任务被 KILL；表预览与长任务抢连接 → 重连后任务重新提交成功", 1),
        ("PySpark round 命名冲突", 0),
        ("Spark 函数与 Python 内置同名互踩 → 别名导入或 F. 前缀", 1),
    ]),
    ("踩坑记录：工具与客户端", [
        ("PyCharm 报“无法解析表” → 元数据缓存未刷新，波浪线 ≠ 执行错误", 0),
        ("Ctrl+Enter 只执行光标所在语句 → 建库语句被跳过（两次中招），逐条核对", 0),
        ("PyCharm DDL 后持续转圈 → 服务端已完成，客户端收尾探测慢，可安全掐掉", 0),
        ("Superset Schema 下拉为空 → 排查 /schemas/ 请求，实为页面入口未进向导", 0),
        ("双击 HTML 报“找不到 chrome” → 默认浏览器指向已卸载程序，重设 .html 关联", 0),
        ("定时任务文件写丢 → 任务计划程序未填“起始于”，文件写进了 System32", 0),
        ("心得：IDE 显示 ≠ 事实，永远去服务端日志找真相", 0),
    ]),
    ("知识掌握地图", [
        ("平台与部署", 0),
        ("WSL2/Linux 运维、tarball 部署、配置文件体系（core/hdfs/mapred/yarn/hive-site）", 1),
        ("数仓工程", 0),
        ("五层分层设计、8 步清洗法、ORC/Parquet 列存、华为 DM 规范、星型维度建模", 1),
        ("SQL 与计算", 0),
        ("窗口函数、谓词下推、惰性执行、Fetch Task 与 MR 任务的区别", 1),
        ("应用开发", 0),
        ("PyHive / pyecharts / Superset / 任务计划程序自动化", 1),
        ("排障方法论", 0),
        ("客户端-服务端分层定位、日志现场还原、基准测试验证、复现-隔离-修复闭环", 1),
    ]),
    ("数据亮点", [
        ("99,457 行真实订单数据，一行不多一行不少（清洗后全链路对账一致）", 0),
        ("5 层数仓 + 2 张维表，全部亲手建表与灌数", 0),
        ("同一清洗逻辑：Hive MR 7 分钟 → Spark 1.5 秒", 0),
        ("列式压缩：6.07MB CSV → 1.39MB ORC（压缩至 23%）", 0),
        ("15+ 个真实故障从定位到修复，全部沉淀成文档", 0),
        ("双可视化路线：Superset 看板 + 每日自动产出的 Python 报表", 0),
    ]),
    ("总结与展望", [
        ("这段经历证明了三件事", 0),
        ("从零能搭：单节点集群、数仓、BI 全链路独立落地", 1),
        ("遇坑能填：每个报错都追到根因，而不是绕过去", 1),
        ("懂取舍：分层、存储格式、引擎选型、工具路线都是权衡题", 1),
        ("下一步方向", 0),
        ("Hive 调优（分桶 / MapJoin / Tez）、实时链路（Kafka + Flink）、集群化部署", 1),
    ]),
]

for title, items in slides:
    s = prs.slides.add_slide(blank)
    add_title(s, title)
    add_bullets(s, items)

prs.save("项目总结.pptx")
print(f"已生成 项目总结.pptx，共 {len(prs.slides.__iter__.__self__._sldIdLst)} 页")
